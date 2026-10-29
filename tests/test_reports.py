"""Pytest Suite for Excel and CSV Report Generation."""

import io
import pytest
from openpyxl import load_workbook

from report_gen import generate_excel_report, generate_csv_report
from config import ExpenseGuardConfig


class TestReportGeneration:
    """Verifies Excel and CSV export fidelity and structure."""

    def test_excel_report_structure(self):
        summary = {
            "total_records": 10,
            "total_amount": 15000.0,
            "passing_count": 8,
            "flagged_count": 2,
            "flagged_amount": 4000.0,
            "passing_amount": 11000.0,
            "compliance_rate": 80.0,
            "duplicates_count": 1,
            "missing_receipts_count": 1,
            "policy_limit_exceeded_count": 0,
            "missing_fields_count": 0,
            "invalid_amounts_count": 0,
            "totals_reconciliation": {
                "calculated_total": 15000.0,
                "claimed_total": 15000.0,
                "mismatch": False,
            },
        }
        items = [
            {
                "transaction_id": "EXP-1",
                "employee_id": "EMP-01",
                "date": "2026-10-10",
                "category": "Meals",
                "description": "Lunch",
                "amount": 2000.0,
                "receipt_attached": True,
                "status": "Needs Review",
                "rule_triggered": "Category Limit Exceeded",
                "evidence": "Exceeds 1500 limit",
                "plain_explanation": "Over limit",
                "recommended_action": "Obtain signoff",
            },
            {
                "transaction_id": "EXP-2",
                "employee_id": "EMP-02",
                "date": "2026-10-10",
                "category": "Travel",
                "description": "Bus",
                "amount": 300.0,
                "receipt_attached": True,
                "status": "Valid",
                "rule_triggered": "None",
                "evidence": "Compliant",
                "plain_explanation": "Valid",
                "recommended_action": "Approve",
            },
        ]
        config = ExpenseGuardConfig()

        excel_bytes = generate_excel_report(summary, items, config)
        assert len(excel_bytes) > 1000

        # Load back workbook
        wb = load_workbook(io.BytesIO(excel_bytes))
        sheet_names = wb.sheetnames
        assert "Executive Summary" in sheet_names
        assert "Review Findings" in sheet_names
        assert "All Transactions" in sheet_names
        assert "Policy Limits" in sheet_names

        # Check Executive Summary content
        ws_sum = wb["Executive Summary"]
        assert ws_sum["A1"].value == "ExpenseGuard AI — Expense Report Audit Dossier"

        # Check Review Findings has 1 flagged row + header
        ws_findings = wb["Review Findings"]
        assert ws_findings.max_row == 2
        assert ws_findings["A2"].value == "EXP-1"

    def test_csv_report_generation(self):
        items = [
            {
                "transaction_id": "EXP-1",
                "employee_id": "EMP-01",
                "date": "2026-10-10",
                "category": "Meals",
                "description": "Lunch",
                "amount": 2000.0,
                "receipt_attached": True,
                "status": "Needs Review",
                "rule_triggered": "Category Limit Exceeded",
                "evidence": "Exceeds 1500 limit",
                "plain_explanation": "Over limit",
                "recommended_action": "Obtain signoff",
            }
        ]
        csv_str = generate_csv_report(items)
        assert "transaction_id,employee_id,date" in csv_str
        assert "EXP-1" in csv_str
        assert "Needs Review" in csv_str
