"""ExpenseGuard AI — Report Generation Module.

Generates professional Excel and CSV audit reports using openpyxl and pandas.
Includes executive summaries, flagged findings, complete transaction ledgers,
policy rules, and human reviewer sign-off blocks.
"""

import io
from typing import Dict, Any, List, Optional
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from config import ExpenseGuardConfig


def generate_excel_report(
    summary: Dict[str, Any],
    items: List[Dict[str, Any]],
    config: Optional[ExpenseGuardConfig] = None,
) -> bytes:
    """Generate a multi-sheet, professionally formatted Excel workbook."""
    wb = Workbook()
    # Remove default sheet
    wb.remove(wb.active)

    # Styling palettes
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    title_font = Font(name="Calibri", size=16, bold=True, color="1E293B")
    subtitle_font = Font(name="Calibri", size=11, italic=True, color="64748B")
    bold_font = Font(name="Calibri", size=11, bold=True)
    normal_font = Font(name="Calibri", size=11)

    primary_fill = PatternFill(start_color="312E81", end_color="312E81", fill_type="solid")  # Indigo-900
    accent_fill = PatternFill(start_color="4338CA", end_color="4338CA", fill_type="solid")   # Indigo-700
    flagged_fill = PatternFill(start_color="FEF3C7", end_color="FEF3C7", fill_type="solid")  # Light Amber
    valid_fill = PatternFill(start_color="D1FAE5", end_color="D1FAE5", fill_type="solid")    # Light Emerald
    header_light_fill = PatternFill(start_color="F1F5F9", end_color="F1F5F9", fill_type="solid")

    thin_border_side = Side(border_style="thin", color="CBD5E1")
    cell_border = Border(left=thin_border_side, right=thin_border_side, top=thin_border_side, bottom=thin_border_side)

    # ==========================================
    # Sheet 1: Executive Summary
    # ==========================================
    ws_summary = wb.create_sheet(title="Executive Summary")
    ws_summary.views.sheetView[0].showGridLines = True

    # Title Banner
    ws_summary["A1"] = "ExpenseGuard AI — Expense Report Audit Dossier"
    ws_summary["A1"].font = title_font
    ws_summary["A2"] = "Automated Compliance & Anomaly Analysis Powered by MCP and Local LLM"
    ws_summary["A2"].font = subtitle_font

    # KPI Table Header
    ws_summary["A4"] = "Audit Metric"
    ws_summary["B4"] = "Value"
    ws_summary["A4"].font = header_font
    ws_summary["A4"].fill = primary_fill
    ws_summary["B4"].font = header_font
    ws_summary["B4"].fill = primary_fill

    kpis = [
        ("Total Transactions Reviewed", f"{summary.get('total_records', 0):,}"),
        ("Total Claimed Amount", f"₹{summary.get('total_amount', 0.0):,.2f}"),
        ("Transactions Passing All Checks", f"{summary.get('passing_count', 0):,}"),
        ("Transactions Flagged for Review", f"{summary.get('flagged_count', 0):,}"),
        ("Total Amount Requiring Review", f"₹{summary.get('flagged_amount', 0.0):,.2f}"),
        ("Total Compliant Amount (Passed Checks)", f"₹{summary.get('passing_amount', 0.0):,.2f}"),
        ("Overall Compliance Rate", f"{summary.get('compliance_rate', 0.0):.1f}%"),
        ("Potential / Exact Duplicates", f"{summary.get('duplicates_count', 0):,}"),
        ("Missing Required Receipts", f"{summary.get('missing_receipts_count', 0):,}"),
        ("Category Policy Limit Breaches", f"{summary.get('policy_limit_exceeded_count', 0):,}"),
        ("Missing Mandatory Fields", f"{summary.get('missing_fields_count', 0):,}"),
        ("Invalid / Negative Amounts", f"{summary.get('invalid_amounts_count', 0):,}"),
    ]

    row_idx = 5
    for label, val in kpis:
        c1 = ws_summary[f"A{row_idx}"]
        c2 = ws_summary[f"B{row_idx}"]
        c1.value = label
        c2.value = val
        c1.font = normal_font
        c2.font = bold_font
        c1.border = cell_border
        c2.border = cell_border
        c2.alignment = Alignment(horizontal="right")
        row_idx += 1

    # Totals Reconciliation Table
    row_idx += 2
    ws_summary[f"A{row_idx}"] = "Totals Reconciliation"
    ws_summary[f"B{row_idx}"] = "Status"
    ws_summary[f"A{row_idx}"].font = header_font
    ws_summary[f"A{row_idx}"].fill = accent_fill
    ws_summary[f"B{row_idx}"].font = header_font
    ws_summary[f"B{row_idx}"].fill = accent_fill

    recon = summary.get("totals_reconciliation", {})
    row_idx += 1
    ws_summary[f"A{row_idx}"] = "Calculated Line Items Sum"
    ws_summary[f"B{row_idx}"] = f"₹{recon.get('calculated_total', summary.get('total_amount', 0.0)):,.2f}"
    ws_summary[f"A{row_idx}"].border = cell_border
    ws_summary[f"B{row_idx}"].border = cell_border
    ws_summary[f"B{row_idx}"].alignment = Alignment(horizontal="right")

    row_idx += 1
    ws_summary[f"A{row_idx}"] = "Claimed Report Total"
    ws_summary[f"B{row_idx}"] = f"₹{recon.get('claimed_total', summary.get('total_amount', 0.0)):,.2f}"
    ws_summary[f"A{row_idx}"].border = cell_border
    ws_summary[f"B{row_idx}"].border = cell_border
    ws_summary[f"B{row_idx}"].alignment = Alignment(horizontal="right")

    row_idx += 1
    ws_summary[f"A{row_idx}"] = "Reconciliation Finding"
    ws_summary[f"B{row_idx}"] = "Discrepancy Detected" if recon.get("mismatch") else "Perfect Match"
    ws_summary[f"A{row_idx}"].border = cell_border
    ws_summary[f"B{row_idx}"].border = cell_border
    ws_summary[f"B{row_idx}"].font = bold_font

    # Human Reviewer Sign-off Section
    row_idx += 3
    ws_summary[f"A{row_idx}"] = "HUMAN FINANCE REVIEWER SIGN-OFF"
    ws_summary[f"A{row_idx}"].font = bold_font
    row_idx += 1
    ws_summary[f"A{row_idx}"] = "Reviewer Name: ___________________________    Date: _______________"
    row_idx += 1
    ws_summary[f"A{row_idx}"] = "Decision:  [  ] Approved as Submitted   [  ] Approved with Exceptions   [  ] Rejected"
    row_idx += 1
    ws_summary[f"A{row_idx}"] = "Reviewer Notes: ____________________________________________________________________"

    # ==========================================
    # Sheet 2: Flagged Findings (Needs Review)
    # ==========================================
    ws_flagged = wb.create_sheet(title="Review Findings")
    ws_flagged.views.sheetView[0].showGridLines = True

    flagged_headers = [
        "Transaction ID",
        "Employee ID",
        "Date",
        "Category",
        "Amount (₹)",
        "Status",
        "Rule Triggered",
        "Supporting Evidence",
        "Plain-Language Explanation",
        "Recommended Next Action",
    ]

    for col_idx, h_text in enumerate(flagged_headers, 1):
        cell = ws_flagged.cell(row=1, column=col_idx, value=h_text)
        cell.font = header_font
        cell.fill = primary_fill
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    flagged_items = [it for it in items if it.get("status") in ["Needs Review", "Invalid Data"]]
    for r_idx, item in enumerate(flagged_items, 2):
        row_vals = [
            item.get("transaction_id"),
            item.get("employee_id"),
            item.get("date"),
            item.get("category"),
            item.get("amount", 0.0),
            item.get("status"),
            item.get("rule_triggered"),
            item.get("evidence"),
            item.get("plain_explanation"),
            item.get("recommended_action"),
        ]
        for c_idx, val in enumerate(row_vals, 1):
            cell = ws_flagged.cell(row=r_idx, column=c_idx, value=val)
            cell.font = normal_font
            cell.border = cell_border
            if c_idx == 5:
                cell.number_format = "₹#,##0.00"
                cell.alignment = Alignment(horizontal="right")
            elif c_idx == 6:
                cell.fill = flagged_fill
                cell.font = bold_font
                cell.alignment = Alignment(horizontal="center")
            elif c_idx in [8, 9, 10]:
                cell.alignment = Alignment(wrap_text=True)

    # ==========================================
    # Sheet 3: All Transactions
    # ==========================================
    ws_all = wb.create_sheet(title="All Transactions")
    ws_all.views.sheetView[0].showGridLines = True

    all_headers = [
        "Transaction ID",
        "Employee ID",
        "Date",
        "Category",
        "Description",
        "Amount (₹)",
        "Receipt Attached",
        "Status",
        "Rule Triggered",
        "Audit Finding Details",
    ]

    for col_idx, h_text in enumerate(all_headers, 1):
        cell = ws_all.cell(row=1, column=col_idx, value=h_text)
        cell.font = header_font
        cell.fill = primary_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")

    for r_idx, item in enumerate(items, 2):
        status = item.get("status", "Passed")
        row_vals = [
            item.get("transaction_id"),
            item.get("employee_id"),
            item.get("date"),
            item.get("category"),
            item.get("description"),
            item.get("amount", 0.0),
            "Yes" if item.get("receipt_attached") else "No",
            status,
            item.get("rule_triggered"),
            item.get("plain_explanation"),
        ]
        for c_idx, val in enumerate(row_vals, 1):
            cell = ws_all.cell(row=r_idx, column=c_idx, value=val)
            cell.font = normal_font
            cell.border = cell_border
            if c_idx == 6:
                cell.number_format = "₹#,##0.00"
                cell.alignment = Alignment(horizontal="right")
            elif c_idx == 8:
                cell.fill = flagged_fill if status in ["Needs Review", "Invalid Data"] else valid_fill
                cell.font = bold_font
                cell.alignment = Alignment(horizontal="center")

    # ==========================================
    # Sheet 4: Policy Thresholds
    # ==========================================
    ws_policy = wb.create_sheet(title="Policy Limits")
    ws_policy.views.sheetView[0].showGridLines = True

    ws_policy["A1"] = "Expense Category"
    ws_policy["B1"] = "Maximum Allowed Limit (₹)"
    ws_policy["A1"].font = header_font
    ws_policy["A1"].fill = primary_fill
    ws_policy["B1"].font = header_font
    ws_policy["B1"].fill = primary_fill

    limits = config.category_limits if config else {}
    p_row = 2
    for cat_name, lim_val in limits.items():
        c1 = ws_policy[f"A{p_row}"]
        c2 = ws_policy[f"B{p_row}"]
        c1.value = cat_name
        c2.value = lim_val
        c1.font = normal_font
        c2.font = bold_font
        c1.border = cell_border
        c2.border = cell_border
        c2.number_format = "₹#,##0.00"
        c2.alignment = Alignment(horizontal="right")
        p_row += 1

    p_row += 1
    c1 = ws_policy[f"A{p_row}"]
    c2 = ws_policy[f"B{p_row}"]
    c1.value = "Receipt Required Threshold"
    c2.value = config.receipt_threshold if config else 500.0
    c1.font = bold_font
    c2.font = bold_font
    c1.border = cell_border
    c2.border = cell_border
    c2.number_format = "₹#,##0.00"

    # Auto-adjust column widths across all sheets
    for ws in [ws_summary, ws_flagged, ws_all, ws_policy]:
        for col in ws.columns:
            max_len = 0
            col_letter = get_column_letter(col[0].column)
            for cell in col:
                val_str = str(cell.value or "")
                # Limit consideration of long paragraphs for width
                if len(val_str) > max_len and len(val_str) < 50:
                    max_len = len(val_str)
            ws.column_dimensions[col_letter].width = max(max_len + 3, 14)

    # Specific overrides for long text columns in Review Findings
    ws_flagged.column_dimensions["G"].width = 24
    ws_flagged.column_dimensions["H"].width = 38
    ws_flagged.column_dimensions["I"].width = 42
    ws_flagged.column_dimensions["J"].width = 38

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return output.getvalue()


def generate_csv_report(items: List[Dict[str, Any]]) -> str:
    """Generate a clean CSV report containing all transactions and audit findings."""
    flat_data = []
    for item in items:
        flat_data.append({
            "transaction_id": item.get("transaction_id"),
            "employee_id": item.get("employee_id"),
            "date": item.get("date"),
            "category": item.get("category"),
            "description": item.get("description"),
            "amount": item.get("amount"),
            "receipt_attached": "Yes" if item.get("receipt_attached") else "No",
            "status": item.get("status"),
            "rule_triggered": item.get("rule_triggered"),
            "evidence": item.get("evidence"),
            "plain_explanation": item.get("plain_explanation"),
            "recommended_action": item.get("recommended_action"),
            "explanation_source": item.get("explanation_source", "deterministic"),
        })

    df = pd.DataFrame(flat_data)
    return df.to_csv(index=False)
